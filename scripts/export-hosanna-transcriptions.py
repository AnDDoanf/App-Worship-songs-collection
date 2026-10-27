"""Export visually transcribed Hosanna lyrics without changing the source workbook."""
import csv
import json
import re
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'exports'
records = {}
for path in OUT.glob('*transcriptions.jsonl'):
    for line in path.read_text(encoding='utf-8-sig').splitlines():
        if not line.strip():
            continue
        record = json.loads(line)
        assert record['id'] not in records, f"Duplicate: {record['id']}"
        records[record['id']] = record
expected = {f'H{n}' for n in range(18, 257)} | {'H129a'}
assert set(records) == expected, (expected - records.keys(), records.keys() - expected)
source = openpyxl.load_workbook(OUT / 'song-collections.xlsx')
sheet = source['HOSANNA']
headers = [cell.value for cell in sheet[1]]
metadata = {}
for cells in sheet.iter_rows(min_row=2):
    row = dict(zip(headers, (cell.value for cell in cells)))
    # The source workbook mistakenly labels Vỗ tay lên as H243.
    if row['id'] == 'H243' and row['songName'] == 'Vỗ tay lên':
        row['id'] = 'H242'
        row['image'] = 'H242_page0'
    if cells[8].hyperlink:
        row['audio'] = cells[8].hyperlink.target
    metadata[row['id']] = row
rows = []
notes = [['H242', 'Corrected source workbook ID H243 and image references to H242 / H242_page0, as printed on the scan of Vỗ tay lên.']]
for song_id in sorted(records, key=lambda s: (int(re.search(r'\d+', s)[0]), s)):
    record = records[song_id]
    row = metadata[song_id].copy()
    row.update({key: value for key, value in record.items() if key in headers})
    lyric = row['lyric']
    assert lyric and lyric.count('[') == lyric.count(']'), song_id
    assert '(Lặp lại)' not in lyric and '[[' not in lyric, song_id
    for image in row['image'].split(','):
        matches = list((ROOT / 'src/assets/images/sheets').glob(f'HosannaVietNam*/{image.strip()}.jpg'))
        assert matches, (song_id, image)
    rows.append([row.get(key) or '' for key in headers])
    if record.get('note'):
        notes.append([song_id, record['note']])

method = [
    'Scope: H18–H256, including H129a; 240 songs from both Hosanna Việt Nam volumes.',
    'Lyrics and inline chords were visually transcribed from the local sheet images. OCR output was not used as the final lyric text.',
    'Lyrics use [chord] notation and literal <br> line breaks for import. No (Lặp lại) labels are included.',
    'Existing workbook metadata and audio links are retained unless a transcription record supplies an override. Metadata and links were not independently verified.',
    'Printed original author credits do not establish authorship of the Vietnamese translation.',
    'Some damaged-font verses remain incomplete: H72, H100, H128, H207, H244. See Review notes for these and other uncertainties or editorial readings.',
]
book = openpyxl.Workbook()
data = book.active
data.title = 'HOSANNA'
data.append(headers)
for row in rows:
    data.append(row)
data.freeze_panes = 'C2'
data.auto_filter.ref = data.dimensions
widths = [10, 36, 10, 15, 40, 25, 30, 24, 45, 110, 35, 24]
for column, width in zip(data.columns, widths):
    data.column_dimensions[column[0].column_letter].width = width
for row in data.iter_rows(min_row=2):
    for cell in row:
        cell.alignment = Alignment(vertical='top', wrap_text=True)
        cell.number_format = '@'
    data.row_dimensions[row[0].row].height = 100
    if str(row[8].value).startswith(('https://', 'http://')):
        row[8].hyperlink = row[8].value
        row[8].style = 'Hyperlink'
review = book.create_sheet('Review notes')
review.append(['id', 'note'])
for note in notes:
    review.append(note)
review.column_dimensions['A'].width = 12
review.column_dimensions['B'].width = 110
review.freeze_panes = 'A2'
info = book.create_sheet('Method')
info.append(['Method and limitations'])
for paragraph in method:
    info.append([paragraph])
info.column_dimensions['A'].width = 120
for ws in book:
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='24476B')
    if ws != data:
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical='top', wrap_text=True)
            ws.row_dimensions[row[0].row].height = 45
base = OUT / 'hosanna-H18-H256'
book.save(base.with_suffix('.xlsx'))
with base.with_suffix('.csv').open('w', encoding='utf-8-sig', newline='') as stream:
    writer = csv.writer(stream)
    writer.writerow(headers)
    writer.writerows(rows)
def markdown_row(values):
    return '| ' + ' | '.join(str(v).replace('|', '&#124;').replace('\n', '<br>') for v in values) + ' |'
base.with_suffix('.md').write_text('\n'.join([markdown_row(headers), markdown_row(['---'] * len(headers)), *(markdown_row(row) for row in rows)]) + '\n', encoding='utf-8')
(OUT / 'hosanna-H18-H256-review.md').write_text('\n\n'.join(method) + '\n\n' + '\n'.join([markdown_row(['id', 'note']), markdown_row(['---', '---']), *(markdown_row(note) for note in notes)]) + '\n', encoding='utf-8')
check = openpyxl.load_workbook(base.with_suffix('.xlsx'), read_only=True)
assert check['HOSANNA'].max_row == 241
assert list(check['HOSANNA'].values)[-1][0] == 'H256'
with base.with_suffix('.csv').open(encoding='utf-8-sig', newline='') as stream:
    assert len(list(csv.reader(stream))) == 241
assert len(base.with_suffix('.md').read_text(encoding='utf-8').splitlines()) == 242
print(f'Exported and validated {len(rows)} songs; {len(notes)} review notes.')
